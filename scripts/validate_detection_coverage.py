#!/usr/bin/env python3
"""Evidence-based validation for the 17 AgenticSOC detection categories.

Each test sends one clearly marked synthetic event, records the database event
UUID returned by SOC Core, and looks up a finding linked to that exact UUID.
A PASS additionally requires the expected deterministic category and the marker
in the finding evidence. Zero-Day is review-only; Cloud-Native and Kubernetes are
reported separately as manual-ingest coverage.
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    import psycopg2
    import psycopg2.extras
except ImportError:  # REST lookup remains fully supported.
    psycopg2 = None


MATRIX_HEADERS = [
    "No", "Kategori Serangan (EN)", "Nama Indonesia", "Log Source / Evidence",
    "Mekanisme Deteksi", "Contoh Rule ID", "MITRE ATT&CK Technique(s)",
    "Status Cakupan", "Metode Validasi", "Hasil Validasi Terakhir", "Catatan / Gap",
]

MATRIX_ROWS = [
    [1, "Malware", "Malware", "Fortigate AV; Wazuh; Defender XDR", "Signature/security alert malware", "RULE-MAL-001", "T1105; T1204", "Automated", "Synthetic malware event; separately test EICAR in an isolated endpoint", "Belum diuji", "Synthetic validation proves normalized rule routing, not AV sensor efficacy"],
    [2, "Ransomware", "Ransomware", "Wazuh FIM + syscheck; Fortigate AV", "Mass rename/encryption behavior", "RULE-RANS-001..003", "T1486; T1490; T1489", "Automated", "Synthetic normalized event; Atomic Red Team T1486 in a non-production host for source validation", "Belum diuji", "Threshold behavior still needs sensor-level validation"],
    [3, "Phishing", "Phishing", "M365 Message Trace; inbox-rule; OAuth consent", "Suspicious inbox rule/OAuth grant", "RULE-PHISH-001..003", "T1566; T1564.008; T1550.001", "Automated", "Synthetic normalized phishing event; separate M365 test tenant scenario", "Belum diuji", "Needs legitimate OAuth baseline"],
    [4, "Credential Theft", "Pencurian Kredensial", "Syslog auth; Wazuh; Entra ID", "Repeated failure/credential attack correlation", "RULE-CRED-001..004", "T1110; T1078; T1552", "Automated", "Synthetic credential_attack event", "Belum diuji", "Correlation thresholds still require source-level validation"],
    [5, "Network Intrusion", "Intrusi Jaringan", "Fortigate IPS; Wazuh network rules", "Confirmed IPS/network attack signal", "RULE-NETI-001..003", "T1190; T1210", "Automated", "Synthetic normalized IPS event; separate lab exploit attempt", "Belum diuji", "Internal category is suspicious_network"],
    [6, "Reconnaissance", "Pengintaian", "Fortigate port-scan detection", "Port-scan/network-scan evidence", "RULE-RECON-001..002", "T1595; T1046", "Automated", "Synthetic reconnaissance event; separate controlled Nmap test", "Belum diuji", "Threshold requires local tuning"],
    [7, "Vulnerability", "Kerentanan", "Wazuh vulnerability detector; CYFIRMA", "CVE match against installed software", "RULE-VULN-001..002", "T1588.006; T1190", "Automated", "Synthetic vulnerability event carrying a reserved test CVE", "Belum diuji", "Provider unavailable must remain unknown, never clean"],
    [8, "File Integrity", "Integritas Berkas", "Wazuh FIM/syscheck", "Critical file checksum change", "RULE-FIM-001..002", "T1565; T1036", "Automated", "Synthetic FIM event; separate isolated-host file change", "Belum diuji", "FIM scan latency needs an SLA"],
    [9, "Supply Chain", "Rantai Pasok", "Wazuh FIM package-manager paths", "Dependency path change", "RULE-SUPPLY-001", "T1195; T1195.001; T1195.002", "Automated", "Synthetic node_modules FIM event", "Belum diuji", "File-level coverage only"],
    [10, "DDoS", "DDoS", "Fortigate DoS sensor", "Volumetric DoS signal", "RULE-DDOS-001", "T1498; T1499", "Automated", "Synthetic dos_attack event; separate controlled lab traffic", "Belum diuji", "Does not prove upstream sensor threshold"],
    [11, "SQL Injection", "Injeksi SQL", "Fortigate WAF/IPS", "SQLi/web-attack signature", "RULE-SQLI-001", "T1190", "Automated", "Synthetic web_attack event; separate non-destructive app test", "Belum diuji", "Signature bypass remains possible"],
    [12, "Insider Threat", "Ancaman Internal", "M365 audit mass download/share", "Bulk activity threshold", "RULE-INSIDER-001", "T1530; T1213", "Automated", "Synthetic insider_risk event", "Belum diuji", "Whitelist backup/migration activity"],
    [13, "Data Exfiltration", "Eksfiltrasi Data", "Microsoft Purview DLP", "DLP policy match", "RULE-DLP-001", "T1041; T1567", "Automated", "Synthetic data_exfiltration event", "Belum diuji", "Depends on sensitivity-label coverage"],
    [14, "APT Activity", "Aktivitas APT", "OTX; ThreatFox IOC attribution", "Threat-actor/IOC attribution", "RULE-APT-001", "T1071; T1105; T1027", "Automated", "Synthetic attributed security alert; live IOC path requires separate provider test", "Belum diuji", "Feed quality and availability remain material"],
    [15, "Zero-Day", "Zero-Day", "Wazuh unfixed vulnerability + CYFIRMA", "Exploit pattern plus no vendor fix", "RULE-ZERODAY-001", "T1190; T1211", "Automated (probabilistic)", "Manual process review; cannot honestly simulate an unknown-unknown", "Tidak dapat diuji otomatis", "Never promise guaranteed zero-day detection"],
    [16, "Cloud-Native", "Serangan Cloud-Native", "Manual POST /events/ingest", "Manual normalized cloud alert", "-", "T1078.004; T1098; T1526", "Manual only", "Manual-ingest pipeline test", "Tidak berlaku", "No native CloudTrail/Azure collector"],
    [17, "Container/Kubernetes", "Serangan Container/K8s", "Manual POST /events/ingest", "Manual normalized K8s alert", "-", "T1610; T1611; T1613", "Manual only", "Manual-ingest pipeline test", "Tidak berlaku", "No native Kubernetes audit collector"],
]


@dataclass(frozen=True)
class CategoryTest:
    name: str
    expected_category: str
    source: str
    event_type: str
    description: str
    severity: str = "high"
    mitre: tuple[str, ...] = ()
    raw: dict[str, Any] = field(default_factory=dict)
    manual_only: bool = False
    applicable: bool = True
    skip_reason: Optional[str] = None

    def payload(self, marker: str) -> dict[str, Any]:
        description = f"[COVERAGE_TEST:{marker}] {self.description}"
        raw = {**self.raw, "test_marker": marker, "is_synthetic_test": True,
               "coverage_expected_category": self.expected_category}
        return {"source": self.source, "type": self.event_type, "description": description,
                "severity": self.severity, "mitre_technique": list(self.mitre),
                "test_marker": marker, "is_synthetic_test": True, "raw": raw}


TESTS: dict[int, CategoryTest] = {
    1: CategoryTest("Malware", "malware", "wazuh", "malware", "Synthetic EICAR-style malware signature alert", mitre=("T1105",), raw={"rule_id": "RULE-MAL-001", "artifact": "EICAR-TEST-SIGNATURE"}),
    2: CategoryTest("Ransomware", "ransomware", "wazuh", "ransomware", "Synthetic mass file encryption and shadow-copy deletion behavior", "critical", ("T1486", "T1490"), {"rule_id": "RULE-RANS-001", "files_renamed_count": 500}),
    3: CategoryTest("Phishing", "phishing", "m365_audit", "phishing", "Synthetic suspicious inbox rule forward-and-delete phishing signal", mitre=("T1566",), raw={"rule_id": "RULE-PHISH-001", "operation": "New-InboxRule"}),
    4: CategoryTest("Credential Theft", "credential_attack", "syslog", "credential_attack", "Synthetic password spray with multiple authentication failures", mitre=("T1110",), raw={"rule_id": "RULE-CRED-001", "failed_attempts": 25}),
    5: CategoryTest("Network Intrusion", "suspicious_network", "fortigate", "network_attack", "Synthetic Fortigate IPS signature attack dropped", mitre=("T1190",), raw={"rule_id": "RULE-NETI-001", "signature": "Exploit.Generic.CoverageTest"}),
    6: CategoryTest("Reconnaissance", "reconnaissance", "fortigate", "reconnaissance", "Synthetic Nmap-like port scan detected", "medium", ("T1595", "T1046"), {"rule_id": "RULE-RECON-001", "ports_scanned": 200}),
    7: CategoryTest("Vulnerability", "vulnerability_management", "wazuh", "vulnerability", "Synthetic CVE-9999-0001 vulnerable package match", mitre=("T1190",), raw={"rule_id": "RULE-VULN-001", "cve": "CVE-9999-0001"}),
    8: CategoryTest("File Integrity", "file_integrity", "wazuh", "fim_change", "FIM: synthetic modification of /etc/coverage_test_critical_file", mitre=("T1565",), raw={"rule_id": "RULE-FIM-001", "fim_path": "/etc/coverage_test_critical_file", "wazuh_rule_groups": ["syscheck"]}),
    9: CategoryTest("Supply Chain", "supply_chain", "wazuh", "fim_change", "FIM: synthetic dependency change in /opt/app/node_modules/coverage-test", "critical", ("T1195.002",), {"rule_id": "RULE-SUPPLY-001", "fim_path": "/opt/app/node_modules/coverage-test", "wazuh_rule_groups": ["syscheck"]}),
    10: CategoryTest("DDoS", "ddos", "fortigate", "dos_attack", "Synthetic volumetric SYN flood DoS sensor signal", "critical", ("T1498",), {"rule_id": "RULE-DDOS-001", "packets_per_second": 500000}),
    11: CategoryTest("SQL Injection", "sql_injection", "fortigate", "web_attack", "Synthetic SQL injection UNION SELECT WAF signature", mitre=("T1190",), raw={"rule_id": "RULE-SQLI-001", "signature": "SQL Injection"}),
    12: CategoryTest("Insider Threat", "insider_threat", "m365_audit", "insider_risk", "Synthetic correlated bulk download by one test account", mitre=("T1530",), raw={"rule_id": "RULE-INSIDER-001", "files_downloaded": 800}),
    13: CategoryTest("Data Exfiltration", "data_exfiltration", "m365_audit", "data_exfiltration", "Synthetic Purview DLP sensitive-data exfiltration policy match", "critical", ("T1041", "T1567"), {"rule_id": "RULE-DLP-001", "operation": "DlpRuleMatch"}),
    14: CategoryTest("APT Activity", "apt_activity", "threat_intel", "security_alert", "Synthetic APT29 threat actor attributed security alert", "critical", ("T1071", "T1105"), {"rule_id": "RULE-APT-001", "attribution": "APT29"}),
    15: CategoryTest("Zero-Day", "zero_day", "review_only", "zero_day", "Not sent", applicable=False, skip_reason="Unknown-unknown cannot be validated honestly with a known signature; perform process and anomaly-detection review."),
    16: CategoryTest("Cloud-Native", "cloud_native", "aws_cloudtrail", "cloud_alert", "Synthetic AWS IAM root UnauthorizedAccess cloud alert", "critical", ("T1078.004", "T1526"), {"cloud_provider": "aws"}, manual_only=True),
    17: CategoryTest("Container/Kubernetes", "container_kubernetes", "k8s_audit", "k8s_alert", "Synthetic privileged container hostPath mount", "high", ("T1610", "T1613"), {"namespace": "coverage-test", "privileged": True}, manual_only=True),
}

STATUS_LABELS = {
    "PASS": "PASS - Finding linked and category verified",
    "PASS_MANUAL_ONLY": "PASS (manual ingest) - Finding linked and category verified",
    "FAIL_NO_FINDING": "FAIL - Event stored but no linked finding",
    "FAIL_WRONG_CATEGORY": "FAIL - Finding category does not match",
    "FAIL_BAD_EVIDENCE": "FAIL - Finding is linked but marker is absent from evidence",
    "ERROR_SEND_FAILED": "ERROR - Event ingestion failed",
    "NOT_TESTABLE_AUTOMATICALLY": "Not testable automatically",
    "DRY_RUN": "Dry-run - No event sent",
}
STATUS_COLORS = {"PASS": "C6EFCE", "PASS_MANUAL_ONLY": "C6EFCE", "NOT_TESTABLE_AUTOMATICALLY": "D9D2E9", "DRY_RUN": "FFEB9C"}


@dataclass
class Config:
    matrix: Path
    output: Path
    report_dir: Path
    soc_core_url: str
    dashboard_url: str
    db_dsn: Optional[str]
    timeout: float
    interval: float
    categories: list[int]
    dry_run: bool


def configure_logging(report_dir: Path) -> logging.Logger:
    report_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("coverage-validation")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
    for handler in (logging.FileHandler(report_dir / "coverage_validation_run.log", encoding="utf-8"), logging.StreamHandler(sys.stdout)):
        handler.setFormatter(formatter); logger.addHandler(handler)
    return logger


def create_matrix(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Coverage Matrix"
    sheet.append(MATRIX_HEADERS)
    for row in MATRIX_ROWS: sheet.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill; cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [7, 24, 24, 38, 38, 24, 38, 22, 52, 28, 56]
    for index, width in enumerate(widths, 1): sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def send_event(cfg: Config, test: CategoryTest, marker: str) -> tuple[Optional[str], str]:
    response = requests.post(f"{cfg.soc_core_url.rstrip('/')}/events/ingest", json=test.payload(marker), timeout=15)
    if response.status_code not in (200, 201, 202):
        return None, f"HTTP {response.status_code}: {response.text[:200]}"
    body = response.json(); event_id = body.get("db_id")
    if not event_id:
        return None, "Ingest response did not contain db_id"
    return str(event_id), f"HTTP {response.status_code}; event stored as {event_id}"


def find_by_api(cfg: Config, event_id: str) -> Optional[dict]:
    response = requests.get(f"{cfg.dashboard_url.rstrip('/')}/api/findings/by-event/{event_id}", timeout=15)
    if response.status_code == 404: return None
    response.raise_for_status()
    return response.json()


def find_by_db(cfg: Config, event_id: str) -> Optional[dict]:
    if not cfg.db_dsn or psycopg2 is None: return None
    with psycopg2.connect(cfg.db_dsn) as connection:
        connection.set_session(readonly=True)
        with connection.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cursor:
            cursor.execute("SELECT * FROM findings WHERE %s::uuid = ANY(event_ids) ORDER BY created_time DESC LIMIT 1", (event_id,))
            row = cursor.fetchone()
            return dict(row) if row else None


def poll_finding(cfg: Config, event_id: str) -> tuple[Optional[dict], float, str]:
    started = time.monotonic(); last_error = ""
    while time.monotonic() - started < cfg.timeout:
        try:
            finding = find_by_db(cfg, event_id) if cfg.db_dsn else find_by_api(cfg, event_id)
            if finding: return finding, round(time.monotonic() - started, 3), "database" if cfg.db_dsn else "dashboard_api"
        except requests.RequestException as exc:
            last_error = f"Finding lookup error: {exc}"
        except Exception as exc:
            last_error = f"Finding lookup error: {type(exc).__name__}: {exc}"
        time.sleep(cfg.interval)
    return None, round(time.monotonic() - started, 3), last_error or ("database" if cfg.db_dsn else "dashboard_api")


def run_test(cfg: Config, number: int, logger: logging.Logger) -> dict[str, Any]:
    test = TESTS[number]; marker = f"{number:02d}-{uuid.uuid4().hex}"
    result: dict[str, Any] = {"category_no": number, "category_name": test.name,
        "expected_category": test.expected_category, "actual_category": None, "test_marker": marker,
        "timestamp_utc": datetime.now(timezone.utc).isoformat(), "manual_only": test.manual_only,
        "coverage_mode": "review only" if not test.applicable else ("manual ingest" if test.manual_only else "normalized automatic path"),
        "event_stored": False, "event_db_id": None, "finding_found": False, "finding_id": None,
        "analysis_status": None,
        "evidence_marker_verified": False, "latency_seconds": None, "lookup_method": None,
        "result_status": None, "note": None, "payload": test.payload(marker) if test.applicable else None}
    if not test.applicable:
        result.update(result_status="NOT_TESTABLE_AUTOMATICALLY", note=test.skip_reason)
        logger.info("[%02d] %s: NOT TESTABLE AUTOMATICALLY", number, test.name); return result
    if cfg.dry_run:
        result.update(result_status="DRY_RUN", note="Payload reviewed; no API request sent")
        logger.info("[%02d] %s: DRY-RUN expected=%s", number, test.name, test.expected_category); return result
    try:
        event_id, note = send_event(cfg, test, marker)
    except (requests.RequestException, ValueError) as exc:
        event_id, note = None, f"Ingest error: {type(exc).__name__}: {exc}"
    if not event_id:
        result.update(result_status="ERROR_SEND_FAILED", note=note)
        logger.error("[%02d] %s: %s", number, test.name, note); return result
    result.update(event_stored=True, event_db_id=event_id, note=note)
    finding, latency, method = poll_finding(cfg, event_id)
    result.update(latency_seconds=latency, lookup_method=method)
    if not finding:
        result.update(result_status="FAIL_NO_FINDING", note=f"{note}; no linked finding within {cfg.timeout:g}s; {method}")
        logger.warning("[%02d] %s: NO FINDING event=%s", number, test.name, event_id); return result
    finding_id = str(finding.get("id") or finding.get("finding_id") or "")
    actual_category = finding.get("category")
    event_ids = [str(value) for value in (finding.get("event_ids") or [])]
    marker_verified = marker in json.dumps(finding, default=str, ensure_ascii=False)
    result.update(finding_found=True, finding_id=finding_id, actual_category=actual_category,
                  analysis_status=finding.get("analysis_status"),
                  evidence_marker_verified=marker_verified)
    if event_id not in event_ids:
        result.update(result_status="FAIL_BAD_EVIDENCE", note="Finding response did not retain the exact event UUID")
    elif actual_category != test.expected_category:
        result.update(result_status="FAIL_WRONG_CATEGORY", note=f"Expected {test.expected_category}, received {actual_category}")
    elif not marker_verified:
        result.update(result_status="FAIL_BAD_EVIDENCE", note="Unique marker absent from finding evidence")
    else:
        result.update(result_status="PASS_MANUAL_ONLY" if test.manual_only else "PASS",
                      note=f"Exact event UUID, expected category, and marker verified via {method}")
    logger.info("[%02d] %s: %s finding=%s category=%s latency=%.3fs", number, test.name,
                result["result_status"], finding_id, actual_category, latency)
    return result


def update_matrix(cfg: Config, results: list[dict[str, Any]]) -> None:
    workbook = load_workbook(cfg.matrix); sheet = workbook["Coverage Matrix"]
    validation_headers = ["Status Validasi Otomatis", "Waktu Uji (UTC)", "Test Marker",
                          "Event UUID (Bukti)", "Finding ID (Bukti)", "Latensi Deteksi (detik)",
                          "Kategori Aktual", "Status Analisis", "Metode Lookup", "Catatan Eksekusi"]
    current = {str(cell.value): cell.column for cell in sheet[1] if cell.value}
    next_column = sheet.max_column + 1
    for header in validation_headers:
        if header not in current: current[header] = next_column; next_column += 1
        cell = sheet.cell(1, current[header], header); cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = 28
    by_number = {item["category_no"]: item for item in results}
    for row in range(2, sheet.max_row + 1):
        item = by_number.get(sheet.cell(row, 1).value)
        if not item: continue
        status = item["result_status"]; label = STATUS_LABELS.get(status, status)
        values = [label, item["timestamp_utc"], item["test_marker"], item["event_db_id"] or "-",
                  item["finding_id"] or "-", item["latency_seconds"] if item["latency_seconds"] is not None else "-",
                  item["actual_category"] or "-", item["analysis_status"] or "-",
                  item["lookup_method"] or "-", item["note"] or "-"]
        for header, value in zip(validation_headers, values):
            cell = sheet.cell(row, current[header], value); cell.alignment = Alignment(wrap_text=True, vertical="top")
        color = STATUS_COLORS.get(status, "FFC7CE")
        sheet.cell(row, current[validation_headers[0]]).fill = PatternFill("solid", fgColor=color)
        sheet.cell(row, 10, label).fill = PatternFill("solid", fgColor=color)
    cfg.output.parent.mkdir(parents=True, exist_ok=True); workbook.save(cfg.output)


def write_reports(cfg: Config, results: list[dict[str, Any]]) -> tuple[Path, Path]:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    json_path = cfg.report_dir / f"coverage_validation_report_{timestamp}.json"
    md_path = cfg.report_dir / f"coverage_validation_report_{timestamp}.md"
    counts: dict[str, int] = {}
    for item in results: counts[item["result_status"]] = counts.get(item["result_status"], 0) + 1
    document = {"run_timestamp_utc": datetime.now(timezone.utc).isoformat(), "mode": "dry-run" if cfg.dry_run else "live",
                "proof_contract": ["SOC Core returned an event database UUID", "Finding event_ids contains that exact UUID",
                                   "Finding category equals expected_category", "Unique marker is present in finding evidence"],
                "counts": counts, "results": results}
    json_path.write_text(json.dumps(document, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    lines = ["# Detection Coverage Validation", "", f"Run (UTC): {document['run_timestamp_utc']}", "",
             "A PASS requires an exact event UUID link, the expected category, and the unique test marker in finding evidence.", "",
             "| No | Category | Coverage mode | Result | Event UUID | Finding ID | Expected → actual | Analysis | Latency |",
             "|---:|---|---|---|---|---|---|---|---:|"]
    for item in results:
        lines.append(f"| {item['category_no']} | {item['category_name']} | {item['coverage_mode']} | "
                     f"{STATUS_LABELS.get(item['result_status'], item['result_status'])} | {item['event_db_id'] or '-'} | "
                     f"{item['finding_id'] or '-'} | {item['expected_category']} → {item['actual_category'] or '-'} | "
                     f"{item['analysis_status'] or '-'} | "
                     f"{item['latency_seconds'] if item['latency_seconds'] is not None else '-'} |")
    lines += ["", "## Counts", ""] + [f"- {STATUS_LABELS.get(key, key)}: **{value}**" for key, value in sorted(counts.items())]
    lines += ["", "## Interpretation", "",
              "This validates normalized ingestion, deterministic finding persistence, and exact event traceability. The analysis status records whether optional AI enrichment completed. It does not prove that Fortigate, Wazuh, M365, Purview, or other source sensors will emit the normalized event in a real attack. Source-level tests remain required for that claim."]
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path, md_path


def parse_args() -> Config:
    parser = argparse.ArgumentParser(description="Validate all AgenticSOC detection categories with traceable evidence.")
    parser.add_argument("--matrix", type=Path, default=Path("Coverage_Matrix_17_Kategori_Serangan.xlsx"))
    parser.add_argument("--output", type=Path)
    parser.add_argument("--report-dir", type=Path, default=Path("reports/coverage"))
    parser.add_argument("--soc-core-url", default="http://localhost:48200")
    parser.add_argument("--dashboard-url", default="http://localhost:38080")
    parser.add_argument("--db-dsn", help="Optional PostgreSQL DSN; API lookup is used when omitted")
    parser.add_argument("--timeout", type=float, default=60)
    parser.add_argument("--interval", type=float, default=2)
    parser.add_argument("--category", type=int, action="append")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--init-matrix", action="store_true", help="Create the baseline 17-category workbook if missing")
    args = parser.parse_args()
    if args.init_matrix and not args.matrix.exists(): create_matrix(args.matrix)
    if not args.matrix.exists(): parser.error(f"matrix not found: {args.matrix} (use --init-matrix to create it)")
    categories = sorted(set(args.category or TESTS.keys()))
    invalid = [number for number in categories if number not in TESTS]
    if invalid: parser.error(f"invalid categories: {invalid}; valid range is 1..17")
    output = args.output or args.matrix.with_name(f"{args.matrix.stem}.validated.xlsx")
    return Config(args.matrix, output, args.report_dir, args.soc_core_url, args.dashboard_url,
                  args.db_dsn, max(1, args.timeout), max(0.2, args.interval), categories, args.dry_run)


def main() -> int:
    cfg = parse_args(); logger = configure_logging(cfg.report_dir)
    logger.info("Starting %s validation for categories %s", "dry-run" if cfg.dry_run else "live", cfg.categories)
    results = [run_test(cfg, number, logger) for number in cfg.categories]
    json_path, md_path = write_reports(cfg, results)
    if not cfg.dry_run: update_matrix(cfg, results)
    logger.info("Reports: %s and %s", json_path, md_path)
    if not cfg.dry_run: logger.info("Validated matrix: %s", cfg.output)
    failing = [item for item in results if item["result_status"].startswith(("FAIL", "ERROR"))]
    return 2 if failing else 0


if __name__ == "__main__":
    raise SystemExit(main())
